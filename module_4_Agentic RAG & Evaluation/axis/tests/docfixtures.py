"""Builders for real, parseable test documents.

At Milestone -1 the acceptance tests uploaded byte strings like
`b"%PDF-1.4 Parental leave is 16 weeks."`. That was fine while ingestion was a
no-op, and stopped being fine the moment a real parser existed: pypdf correctly
refuses it, so the tests would have been asserting against a document that never
indexed. Fixtures have to be genuine files or the criteria they check are
meaningless.

Built here rather than committed as binaries because a `.pdf` in the repo is
opaque — nobody can review what a test document says in a diff, and a fixture
whose content nobody can read is a fixture nobody trusts.

The Office formats are written by the same libraries that read them
(`python-pptx`, `python-docx`, `openpyxl`), so a fixture exercises the real
round trip. Only the PDF is assembled by hand — no reportlab needed for a
single page of text.
"""

from __future__ import annotations

import io


def make_pdf(*pages: str) -> bytes:
    """A minimal but genuinely valid PDF, one page per argument.

    Assembled by hand: each page carries a content stream that *shows text*
    (`BT /F1 12 Tf (…) Tj ET`), which is what makes the text extractable. A PDF
    with pages but no text-showing operators extracts to nothing — the exact trap
    the old `b"%PDF-1.4 …"` fixtures fell into, and why `parsers.py` treats "no
    readable text" as a failure rather than an empty success.

    Offsets in the xref table are computed from the real byte positions, because
    pypdf validates them.
    """
    if not pages:
        pages = ("",)

    objects: list[bytes] = []

    # 1: catalog, 2: page tree, 3: font. Pages and their streams follow in pairs.
    page_ids = [4 + 2 * i for i in range(len(pages))]
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)

    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode("latin-1")
    )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for index, text in enumerate(pages):
        content_id = page_ids[index] + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 3 0 R >> >> "
                f"/Contents {content_id} 0 R >>"
            ).encode("latin-1")
        )

        # One Tj per line, moved down the page by TL/T*. Parentheses and
        # backslashes are PDF string delimiters and must be escaped.
        lines = text.split("\n") or [""]
        shown = "\n".join(f"({_escape(line)}) Tj T*" for line in lines)
        stream = f"BT /F1 12 Tf 14 TL 72 720 Td\n{shown}\nET".encode("latin-1")
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1")
            + stream
            + b"\nendstream"
        )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    # A binary comment line marks the file as containing binary data; some readers
    # rely on it and it costs four bytes.
    out.write(b"%\xe2\xe3\xcf\xd3\n")

    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode("latin-1"))
        out.write(body)
        out.write(b"\nendobj\n")

    xref_offset = out.tell()
    count = len(objects) + 1
    out.write(f"xref\n0 {count}\n".encode("latin-1"))
    # Entry 0 is always the head of the free list, in this exact format.
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode("latin-1"))
    out.write(
        f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode(
            "latin-1"
        )
    )
    return out.getvalue()


def _escape(text: str) -> str:
    return (
        text.replace("\\", r"\\")
        .replace("(", r"\(")
        .replace(")", r"\)")
        # PDF simple strings are Latin-1; anything outside it would corrupt the
        # stream, so it is dropped rather than silently mangled.
        .encode("latin-1", "ignore")
        .decode("latin-1")
    )


def make_pptx(*slides: tuple[str, str]) -> bytes:
    """A real .pptx. Each slide is a (title, body) pair."""
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]  # Title and Content
    for title, body in slides:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def make_docx(
    *sections: tuple[str, str],
    table: list[list[str]] | None = None,
) -> bytes:
    """A real .docx. Each section is a (heading, body) pair.

    `table` is appended under the last heading, so tests can check that a table
    stays attached to the section it belongs to rather than drifting to the end.
    """
    from docx import Document

    document = Document()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, value in enumerate(row):
                grid.cell(row_index, col_index).text = str(value)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    """A real .xlsx. Each sheet is a list of rows."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_png(width: int = 64, height: int = 64, colour: str = "navy") -> bytes:
    """A real PNG. Content is irrelevant — captioning is faked in tests."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(buffer, format="PNG")
    return buffer.getvalue()


def make_zip_bomb(*, member_size: int = 50 * 1024 * 1024, members: int = 20) -> bytes:
    """A small .xlsx-named archive that claims to expand to ~1 GB.

    Highly compressible zeroes, so the file on disk is tiny while its directory
    advertises an enormous uncompressed size — which is precisely what
    `ingestion/safety.py` reads to refuse it without decompressing anything.
    """
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for index in range(members):
            archive.writestr(f"xl/worksheets/sheet{index}.xml", b"\0" * member_size)
    return buffer.getvalue()
