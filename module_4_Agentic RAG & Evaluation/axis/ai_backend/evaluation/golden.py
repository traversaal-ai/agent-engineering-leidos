"""Loading the golden Q&A set and the corpus it is measured against.

**The corpus is `data/` at the repo root, and the ground truth is here.** The two used
to sit together under `golden/fixtures/`, which filed the corpus as test data — and it
is not: those five files are what a class loads with one click and asks questions of.
`data/` is the product asset; `questions.yaml` beside this module is the *measurements
about* it. Splitting them is what makes each name honest.

The five ACME Aerospace documents were **copied from** `reference/module_3_Enterprise
RAG/data/` — the corpus this cohort was already taught on — with one edit, recorded in
`demo_documents()` below and in `data/README.md`. `reference/` is read by people, never
by Axis: nothing in this codebase opens a path inside it. They replaced a synthetic
handbook-and-security-policy pair, and the reason was not continuity alone: the ACME
documents *interlock*, so a question can require using what one says to know what to
look for in another.

They are committed as plain text so anyone can read in a diff what the evaluation
is actually asking about — a `.pdf` fixture is opaque, and a fixture nobody can
read is a fixture nobody can debug against.

The table and image fixtures are *generated* rather than committed, for the same
reason: a checked-in `.xlsx` is a binary blob, whereas the code below states its
contents in eight readable lines. Generating them also means the real `openpyxl`
and `Pillow` paths get exercised, so a parser regression shows up here rather
than only in production.
"""

from __future__ import annotations

import io
from pathlib import Path

import yaml
from pydantic import BaseModel, Field

# The corpus, at the repo root — two levels up from `ai_backend/evaluation/`. Same
# `__file__`-relative idiom as `CACHE_DIR` in `providers/search.py`, rather than anything
# relative to the working directory: `python -m axis` can be run from anywhere.
#
# **No fallback if it is missing.** A loader that quietly dropped back to a second
# location would leave the demo button reporting "0 indexed" in front of a class, which is
# what `CachedSearchProvider` refuses to do by construction and for the same reason.
# `test_the_demo_corpus_is_on_disk_where_the_loader_looks` makes a wrong path a red suite
# instead.
CORPUS_DIR = Path(__file__).resolve().parents[2] / "data"
# The one `.md` in the corpus directory that is *about* the corpus rather than part of
# it. See the exclusion in `demo_documents()`.
CORPUS_README = "README.md"
GOLDEN_DIR = Path(__file__).parent / "golden"
QUESTIONS_FILE = GOLDEN_DIR / "questions.yaml"


class GoldenQuestion(BaseModel):
    id: str
    question: str
    kind: str
    expect_documents: list[str] = Field(default_factory=list)
    expect_contains: list[str] = Field(default_factory=list)
    expect_grounded: bool = True
    # The hand-authored ground-truth decomposition, for questions that have one.
    #
    # Two jobs, and the first is why it exists. Retrieving over these instead of
    # over the whole question measures *the value of decomposition* with the
    # model's ability to decompose held out — which is the only way to see the
    # agentic win offline, where a fake LLM cannot split anything (System Design
    # Section 12). On a real-provider run it is also the ground truth the
    # decomposer's own `sub_questions_text` can be scored against.
    #
    # Empty for a question one lookup answers. Deliberately not defaulted to
    # `[question]`: "nobody authored a decomposition" and "the correct
    # decomposition is the question itself" are different facts, and only the
    # first should be silent.
    sub_questions: list[str] = Field(default_factory=list)

    # The ordered chain a working agent should follow, where each step is knowable
    # only from the previous step's result.
    #
    # **Not interchangeable with `sub_questions`, and the distinction is the
    # lesson.** A compound question's parts are all visible in its wording, so a
    # decomposer can split it; a multi-hop question's second part is *not* — the
    # words needed to search for it appear only in the first hop's answer. Held in
    # its own field so the two mechanisms are measured separately: decomposition
    # cannot help here, and a set that conflated them would credit it with a win
    # it did not earn.
    hops: list[str] = Field(default_factory=list)

    # A follow-up needs a turn in front of it to mean anything. `context_question`
    # is that turn; `resolved` is this question with its references filled in, as a
    # working rewriter should produce it. The delta between retrieving for
    # `question` and retrieving for `resolved` is the value of memory, with the
    # model's ability to resolve held out.
    context_question: str = ""
    resolved: str = ""

    # This question is about the whole of its documents rather than a passage in
    # them, so what matters is how much of them was read.
    expect_full_coverage: bool = False

    # The facts that must be **retrieved** — one per side of a compound question, one
    # per hop of a chain.
    #
    # **Distinct from `expect_contains`, which is what must be *stated*.** That one
    # gates `answer_passed` and requires every phrase to appear in the model's prose;
    # overloading it would silently make every answer test stricter. These are about
    # what came back from the index, whether or not the model then used it.
    #
    # **Why they exist at all.** `document_recall` is `len(expected & retrieved) /
    # len(expected)`, and every compound question here names exactly two documents — so
    # it is a three-valued metric that one lucky chunk per document scores 1.00 on. It
    # cannot tell "reached both figures being compared" from "reached one of them and
    # something else in the other file", and under a real embedding model that is
    # precisely the difference decomposition buys. `scoring.py`'s own docstring names
    # the blindness; this is the ground truth that lets a measure see past it.
    expect_passages: list[str] = Field(default_factory=list)

    @property
    def is_unanswerable(self) -> bool:
        return not self.expect_grounded

    @property
    def is_compound(self) -> bool:
        """Whether splitting this question is expected to change what it finds."""
        return len(self.sub_questions) > 1

    @property
    def is_multi_hop(self) -> bool:
        """Whether answering needs a lookup the question itself cannot name."""
        return len(self.hops) > 1

    @property
    def is_follow_up(self) -> bool:
        """Whether this question is meaningless without the turn before it."""
        return bool(self.context_question and self.resolved)


class GoldenDocument(BaseModel):
    """A fixture document, named by the stem the questions refer to."""

    name: str
    filename: str
    data: bytes


class GoldenSet(BaseModel):
    questions: list[GoldenQuestion]
    documents: list[GoldenDocument]

    def by_kind(self, kind: str) -> list[GoldenQuestion]:
        return [q for q in self.questions if q.kind == kind]

    @property
    def answerable(self) -> list[GoldenQuestion]:
        return [q for q in self.questions if q.expect_grounded]

    @property
    def unanswerable(self) -> list[GoldenQuestion]:
        return [q for q in self.questions if not q.expect_grounded]


def load_questions(path: Path | None = None) -> list[GoldenQuestion]:
    raw = yaml.safe_load((path or QUESTIONS_FILE).read_text(encoding="utf-8"))
    return [GoldenQuestion(**entry) for entry in raw["questions"]]


def load_documents() -> list[GoldenDocument]:
    """Every fixture: the committed Markdown plus the generated table and image.

    **This is the harness corpus, and it is deliberately larger than the demo
    corpus.** The two used to be one function, which quietly coupled them: the
    demo endpoint is bound by `max_files_per_session` (five), so anything added
    here for measurement had to fit inside a limit that has nothing to do with
    measurement. The harness ingests through the runtime directly and is bound by
    no such thing.

    The split is what lets both be right. Section 12 of the System Design asks for
    coverage per document *type*, so the generated `.xlsx` and `.png` stay here and
    keep the tabular and visual questions measurable. `demo_document_set()` in
    `backend/dispatch.py` offers the five ACME documents alone — which is what the
    class has already seen, and exactly the file limit.
    """
    documents = list(demo_documents())
    documents.append(
        GoldenDocument(name="rates", filename="rates.xlsx", data=_build_rates_xlsx())
    )
    documents.append(
        GoldenDocument(
            name="headcount", filename="headcount.png", data=_build_headcount_png()
        )
    )
    return documents


def demo_documents() -> list[GoldenDocument]:
    """The corpus a class loads with one click: the five ACME documents.

    They live in `data/` at the repo root — see `data/README.md`. Copied from
    `reference/module_3_Enterprise RAG/data/`, which is the material this cohort has
    already worked with, so a question asked here lands on documents a student recognises
    rather than on a synthetic handbook nobody has read. The copy in `data/` is the
    authoritative one; `reference/` is the record of the demo they were shown and is not
    read at runtime.

    They also do something the synthetic fixtures could not: **they interlock.** A
    slipped deliverable appears as an action item in the kickoff notes, an
    escalation in the status review, and risk R-02 in the register, so a question
    can genuinely require using what one document says to know what to look for in
    another. That is the multi-hop pain point, and without a corpus that has a real
    chain in it there is nothing to demonstrate.

    One edit on the way in, and it is applied to the **files** rather than here: the
    three-line "FICTIONAL SAMPLE DOCUMENT" blockquote that opens all five is collapsed
    to a single line. Module 3 strips it entirely before indexing and records why —
    identical across five files, it made every document's first chunk look alike and
    outrank real content. Keeping one line keeps the provenance a fictional corpus needs
    without five copies of it competing for the top of every search.

    **Editing the files rather than the loader is deliberate.** The guarantee worth
    having on a teaching tool is the simplest one — the bytes on disk are the bytes that
    get indexed — and a loader that rewrote documents on the way in would be a
    transformation a student cannot see, on the one screen whose whole claim is that they
    can see everything. It also means dropping your own file into `data/` indexes it
    exactly as written. The cost is that the edit is only a committed diff, which is how
    it came to be mistaken for a path problem; `data/README.md` and
    `test_the_provenance_notice_is_one_line` are what pay for that.
    """
    return [
        GoldenDocument(name=path.stem, filename=path.name, data=path.read_bytes())
        for path in sorted(CORPUS_DIR.glob("*.md"))
        if path.name != CORPUS_README
        # `data/README.md` documents the corpus; it is not part of it. Excluded by name
        # rather than by moving the corpus into a subdirectory, because one skipped
        # filename is cheaper to understand than a level of nesting — and cheaper than
        # the alternative it replaced, which was a directory that could hold no notes at
        # all. Without this the README indexes as a sixth document and the demo endpoint
        # starts returning 413 against a five-file limit.
    ]


def load_golden_set(path: Path | None = None) -> GoldenSet:
    return GoldenSet(questions=load_questions(path), documents=load_documents())


def _build_rates_xlsx() -> bytes:
    """The tabular fixture. Contents stated here so they are reviewable."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    for row in (
        ("Band", "Daily rate EUR", "Notice period"),
        ("Junior", 450, "4 weeks"),
        ("Mid", 700, "6 weeks"),
        ("Senior", 950, "8 weeks"),
        ("Principal", 1200, "12 weeks"),
    ):
        sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_headcount_png() -> bytes:
    """The visual fixture — a simple bar chart with labels.

    Drawn rather than committed, and deliberately *legible*: the caption a vision
    model produces is only useful if there is real text in the image to
    transcribe. A blank rectangle would test the plumbing while telling us nothing
    about whether captioning helps retrieval.
    """
    from PIL import Image, ImageDraw

    width, height = 480, 300
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((12, 10), "Headcount by team, 2026", fill="black")

    bars = [("Platform", 24), ("Product", 18), ("Data", 11), ("Design", 7)]
    scale = 8
    for index, (team, count) in enumerate(bars):
        top = 50 + index * 55
        draw.rectangle([110, top, 110 + count * scale, top + 30], fill="navy")
        draw.text((12, top + 10), team, fill="black")
        draw.text((118 + count * scale, top + 10), str(count), fill="black")

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()
